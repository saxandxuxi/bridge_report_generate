# -*- coding: utf-8 -*-
"""冒烟测试：覆盖近期修复的高风险路径。

运行：python -m unittest tests.test_smoke
（已同步到 requirements.txt 的 pytest 也可直接 pytest tests/）
"""
import datetime as dt
import importlib.util
import json
import os
import re
import tempfile
import unittest

from docx import Document

from report_agent.report_builder import (
    _apply_period_text_fixes,
    _expand_row_tables,
    _unique_cells,
    build_value_resolver,
    format_report_number,
    verify_table_columns,
)
from report_agent.template_analyzer import analyze_template
from run_placeholder_audit import (
    _is_obviously_static,
    _key_valid,
    _options_for,
)
from report_agent.placeholder_arbiter import (
    arbitrate_numbers,
    apply_verdict,
    build_options,
    is_obviously_static,
)
from report_agent.bridge_source import BridgeData
from report_agent.reviewer import self_check_report


class RowsExpansionTest(unittest.TestCase):
    def test_col_resolver_accepts_table_kwargs(self):
        """{{rows.*}} 展开时 _fill_paragraph 传入的
        table_title/row_index 关键字不得导致 TypeError。"""
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "t.docx")
            doc = Document()
            t = doc.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text = "{{rows.daily_records}}"
            t.rows[0].cells[1].text = "{{col.date}}/{{col.value}}"
            doc.save(p)

            doc = Document(p)
            datasets = {
                "daily_records": [
                    {"date": "2026-01-01", "value": "3.5"},
                    {"date": "2026-01-02", "value": "4.2"},
                ]
            }
            n = _expand_row_tables(
                doc, datasets, lambda key, **kw: "")
            self.assertEqual(n, 2)
            texts = [c.text for r in doc.tables[0].rows
                     for c in r.cells]
            joined = "".join(texts)
            self.assertIn("2026-01-01/3.5", joined)
            self.assertIn("2026-01-02/4.2", joined)


class PeriodFixTest(unittest.TestCase):
    def test_stale_year_month_range(self):
        doc = Document()
        doc.add_paragraph("综上，大桥在2025年01月-03月之间正常。")
        n = _apply_period_text_fixes(doc, {
            "start": dt.date(2026, 1, 1),
            "end": dt.date(2026, 3, 31),
            "mode": "quarterly",
        })
        self.assertEqual(n, 1)
        self.assertIn("2026年1月至3月之间", doc.paragraphs[0].text)


class MergedCellTest(unittest.TestCase):
    def test_unique_cells_dedupe(self):
        doc = Document()
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).merge(t.cell(0, 1))
        raw = len(t.rows[0].cells)
        unique = len(_unique_cells(t.rows[0]))
        self.assertLess(unique, raw)


class NumberScanTest(unittest.TestCase):
    def test_mixed_paragraph_numbers_found(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "t.docx")
            doc = Document()
            doc.add_paragraph("year {{date.period_label_cn}} total 3.5 m")
            doc.save(p)
            result = analyze_template(p)
            nums = [x["number"] for x in result["candidate_numbers"]]
            self.assertIn("3.5", nums)
            self.assertNotIn("{{date.period_label_cn}}", nums)


class VerifyTableTest(unittest.TestCase):
    def test_same_value_column_reports_sensor_count(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "out.docx")
            doc = Document()
            t = doc.add_table(rows=3, cols=2)
            t.rows[0].cells[0].text = "测点"
            t.rows[0].cells[1].text = "值"
            t.rows[1].cells[0].text = "A"
            t.rows[1].cells[1].text = "12.34"
            t.rows[2].cells[0].text = "B"
            t.rows[2].cells[1].text = "12.34"
            doc.save(p)
            lineage = [
                {"占位符": "cell.x.A.avg", "输出": "12.34",
                 "传感器": {"传感器编号": "111"}},
                {"占位符": "cell.x.B.avg", "输出": "12.34",
                 "传感器": {"传感器编号": "222"}},
            ]
            warns = verify_table_columns(p, lineage, label="t")
            same = [w for w in warns if w.get("问题") == "整列同值"]
            self.assertTrue(same)
            self.assertEqual(same[0]["血缘命中传感器数"], 2)


class AuditHelpersTest(unittest.TestCase):
    def test_key_valid_schema(self):
        mkeys = {"temperature", "wind_speed", "structure_temperature",
                 "displacement"}
        self.assertTrue(_key_valid("stats.temperature.max", mkeys))
        self.assertTrue(_key_valid("stats.wind_speed.avg", mkeys))
        # 关键词表里没有的统计量 / 非法结构一律拒绝
        self.assertFalse(_key_valid("stats.wind.10min_max", mkeys))
        self.assertFalse(_key_valid(
            "cell.temperature.time_series.section_1_2", mkeys))
        self.assertFalse(_key_valid("stats.bridge.width.avg", mkeys))

    def test_options_only_from_vocab(self):
        opts = _options_for({
            "text": "最大10min平均风速为11.93 m/s。",
            "rule_placeholder": None,
        })
        self.assertIn("stats.wind_speed.max", opts)
        self.assertTrue(all(o.startswith("stats.") for o in opts))

    def test_chart_id_and_citation_are_static(self):
        self.assertTrue(_is_obviously_static(
            {"text": "1_2主跨钢桁梁桥面板_应变_histogram"}))
        self.assertTrue(_is_obviously_static(
            {"text": "《公路桥梁…》JTG/T D65-05-2015"}))
        self.assertFalse(_is_obviously_static(
            {"text": "结构温度最大值为40.59℃"}))
        # 同段有规范引用，但真实数据离引用较远 -> 不能判静态
        mixed = "依据《公路桥梁结构监测技术规范》(JT/T 1037-2022)，最高湿度为100.55％。"
        self.assertFalse(is_obviously_static(
            mixed, mixed.index("100.55")))


class ArbiterTest(unittest.TestCase):
    def test_build_options_and_static(self):
        opts = build_options("最大10min平均风速为11.93 m/s。")
        self.assertIn("stats.wind_speed.max", opts)
        self.assertTrue(is_obviously_static("…_应变_histogram"))
        self.assertTrue(is_obviously_static("JTG/T D65-05-2015"))

    def test_apply_verdict_flip_and_key(self):
        n = {"verdict": "replace", "placeholder": None, "reasons": []}
        action = apply_verdict(n, "keep", None, ["stats.temperature.max"],
                               confidence=0.9)
        self.assertEqual(action, "flip_keep")
        self.assertEqual(n["verdict"], "keep")
        # 低置信不翻案
        n2 = {"verdict": "replace", "placeholder": None, "reasons": []}
        self.assertEqual(
            apply_verdict(n2, "keep", None, [], confidence=0.5),
            "keep_rule")
        self.assertEqual(n2["verdict"], "replace")
        # key 必须来自 options
        n3 = {"verdict": "review", "placeholder": None, "reasons": []}
        self.assertEqual(
            apply_verdict(n3, "replace", "stats.temperature.max",
                          ["stats.wind_speed.max"], 0.9),
            "replace_no_key")
        self.assertIsNone(n3["placeholder"])
        n4 = {"verdict": "review", "placeholder": None, "reasons": []}
        self.assertEqual(
            apply_verdict(n4, "replace", "stats.wind_speed.max",
                          ["stats.wind_speed.max"], 0.9),
            "set_key")
        self.assertEqual(n4["placeholder"], "stats.wind_speed.max")


class ChartLayoutTest(unittest.TestCase):
    """总标题与子图标题、图例与坐标轴不得重叠。"""

    def test_suptitle_does_not_overlap_axes_titles(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        spec = importlib.util.spec_from_file_location(
            "bcl_layout",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))),
                "preprocess", "scripts", "build_chart_library.py"))
        bcl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(bcl)
        fig, axes = plt.subplots(2, 2, figsize=(10, 8))
        for i, ax in enumerate(axes.reshape(-1)):
            ax.plot([0, 1], [0, 1])
            ax.set_title(f"子图{i}", fontsize=12)
        fig.suptitle("总标题", fontsize=16)
        bcl._finalize_layout(fig, rows=2)
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        su = fig._suptitle.get_window_extent(r)

        def _ov(a, b):
            return not (a.x1 <= b.x0 or b.x1 <= a.x0
                        or a.y1 <= b.y0 or b.y1 <= a.y0)
        for ax in axes.reshape(-1):
            self.assertFalse(_ov(su, ax.title.get_window_extent(r)))
        plt.close(fig)


class ConclusionsFallbackTest(unittest.TestCase):
    """{{conclusions}} 不得因桥数据缺失/异常而中断报告。"""

    def _period(self):
        return {"start": dt.date(2026, 1, 1),
                "end": dt.date(2026, 3, 31),
                "label": "2026.1~3",
                "label_cn": "2026年第1季度"}

    def test_no_bridge_returns_marker(self):
        sink = []
        r = build_value_resolver({}, self._period(), bridge=None,
                                 missing_sink=sink)
        self.assertEqual(r("conclusions"), "—")
        self.assertIn("conclusions", sink)

    def test_bridge_exception_returns_marker(self):
        class BadBridge:
            def build_conclusions(self, period, llm_cfg=None):
                raise RuntimeError("boom")

        sink = []
        r = build_value_resolver({}, self._period(), bridge=BadBridge(),
                                 missing_sink=sink)
        self.assertEqual(r("conclusions"), "—")

    def test_bridge_text_returned(self):
        class OkBridge:
            def build_conclusions(self, period, llm_cfg=None):
                return "（1）桥梁结构处于良好状态。"

        r = build_value_resolver({}, self._period(), bridge=OkBridge())
        self.assertIn("良好状态", r("conclusions"))


class ReportNumberFormatTest(unittest.TestCase):
    def test_number_rules(self):
        self.assertEqual(format_report_number(21.5665), "21.57")
        self.assertEqual(format_report_number(-3.371), "-3.37")
        self.assertEqual(format_report_number(-15.123), "-15.1")
        self.assertEqual(format_report_number(0.001234), "0.001234")
        self.assertEqual(format_report_number(0.000159), "0.000159")
        self.assertEqual(format_report_number(1.234e-7), "1.234e-07")
        self.assertEqual(format_report_number(0), "0")
        self.assertEqual(format_report_number(90), "90")
        self.assertEqual(format_report_number(float("nan")), "—")
        self.assertEqual(format_report_number(float("inf")), "—")


class GrossStatFaultTest(unittest.TestCase):
    def test_vibration_extreme_only_invalidates_extremes(self):
        st = {"最大值": 5602250.0, "最小值": -0.0975,
              "平均值": 5.0e-4, "差值": 5602250.1}
        f = "DZJSD(yJsd)"
        self.assertTrue(BridgeData._gross_stat_fault(st, f, "max"))
        self.assertTrue(BridgeData._gross_stat_fault(st, f, "range"))
        self.assertTrue(BridgeData._gross_stat_fault(st, f, "abs_max"))
        self.assertFalse(BridgeData._gross_stat_fault(st, f, "min"))
        self.assertFalse(BridgeData._gross_stat_fault(st, f, "avg"))

    def test_strain_extreme(self):
        st = {"最大值": 1.2e6, "最小值": -88.0, "平均值": 112.0,
              "差值": 1.2e6}
        self.assertTrue(BridgeData._gross_stat_fault(st, "YB(rsg)", "max"))
        self.assertFalse(BridgeData._gross_stat_fault(st, "YB(rsg)", "min"))

    def test_temp_summer_anomaly(self):
        st = {"最大值": 30.25, "最小值": -30.7, "平均值": 0.0,
              "差值": 60.95}
        self.assertTrue(BridgeData._gross_stat_fault(
            st, "WD(temp)", "min"))
        self.assertTrue(BridgeData._gross_stat_fault(
            st, "WD(temp)", "range"))

class StatDirectReadTest(unittest.TestCase):
    """统计库有什么就填什么，没有就填“—”，报告不再二次判断。"""

    def _bridge(self, tmp, stats_entry, sensor_id="636"):
        pos = "汝城侧中跨1/4截面底板上游"
        safe = re.sub(r'[\\/:*?"<>|]', "_", pos).strip()
        d = os.path.join(tmp, "位置统计", safe)
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, "WD(temp).json"), "w",
                  encoding="utf-8") as f:
            json.dump(stats_entry, f, ensure_ascii=False)
        cfg = dict(json.load(open(
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))), "config",
                "config_mishuihe.json"), encoding="utf-8"))["bridge_data"])
        cfg["stats_dir"] = tmp
        cfg["charts_dir"] = tmp
        b = BridgeData(cfg)
        b.load()
        return b

    @staticmethod
    def _entry(**stats):
        base = {"起始日期": "2026-07-01", "结束日期": "2026-09-30",
                "覆盖天数": 90}
        base.update(stats)
        return {"汝城侧中跨1/4截面底板上游": {
            "测点1": {"统计": base, "传感器编号": "636",
                      "特征": "WD(temp)"}}}

    def test_values_are_read_verbatim(self):
        with tempfile.TemporaryDirectory() as tmp:
            b = self._bridge(tmp, self._entry(平均值=19.066287, 最大值=20.75,
                                              最小值=0.0, 差值=20.75))
            period = {"start": dt.date(2026, 7, 1),
                      "end": dt.date(2026, 9, 30), "label": "2026.7~9"}
            got = {s: b._sensor_stat("636", "structure_temperature", s, period,
                                     feature="WD(temp)")
                   for s in ("平均温度", "最高温度", "最低温度", "最大温差")}
            # 库里的原值原样回填（不做“极值可疑就改写/借别的测点”）
            self.assertAlmostEqual(got["平均温度"], 19.066287, places=6)
            self.assertAlmostEqual(got["最高温度"], 20.75, places=6)
            self.assertAlmostEqual(got["最低温度"], 0.0, places=6)
            self.assertAlmostEqual(got["最大温差"], 20.75, places=6)

    def test_missing_stat_returns_none(self):
        with tempfile.TemporaryDirectory() as tmp:
            b = self._bridge(tmp, self._entry(平均值=19.0, 最大值=20.0))
            period = {"start": dt.date(2026, 7, 1),
                      "end": dt.date(2026, 9, 30), "label": "2026.7~9"}
            self.assertIsNone(b._sensor_stat("636", "structure_temperature",
                                             "最小值", period,
                                             feature="WD(temp)"))
            self.assertAlmostEqual(
                b._sensor_stat("636", "structure_temperature", "平均值",
                               period, feature="WD(temp)"), 19.0, places=6)

    def test_stored_value_wins_over_daily_recompute(self):
        """整体统计里有的字段直接回填，不再用每日明细“重算清洗”。"""
        with tempfile.TemporaryDirectory() as tmp:
            b = self._bridge(tmp, self._entry(
                平均值=5.0e-4, 最大值=5602250.0, 最小值=-0.0975,
                差值=5602250.1,
                每日统计=[{"日期": "2026-07-01", "最大值": 4.75,
                           "最小值": -1.64, "平均值": 0.0}]))
            period = {"start": dt.date(2026, 7, 1),
                      "end": dt.date(2026, 9, 30), "label": "2026.7~9"}
            v = b._sensor_stat("636", "structure_temperature", "max", period,
                               feature="WD(temp)")
            self.assertAlmostEqual(v, 5602250.0, places=1)


class SelfCheckPhysicsTest(unittest.TestCase):
    def test_physical_and_placeholder_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "r.docx")
            doc = Document()
            t = doc.add_table(rows=2, cols=1)
            t.rows[0].cells[0].text = "最大值(m/s²)"
            t.rows[1].cells[0].text = "5602250"
            doc.add_paragraph("strain_左幅3#墩根部顶板_scatter_59")
            doc.save(p)
            issues = self_check_report(p)
            types = {i["type"] for i in issues}
            self.assertIn("physical_range", types)
            self.assertIn("chart_placeholder", types)


class CategoryIsolationTest(unittest.TestCase):
    """总结极值只允许在该指标类别的传感器内聚合，不共享特征码回退。"""

    def _bridge(self):
        b = BridgeData({})
        b.metric_category = {"vibration": "振动",
                             "earthquake_load": "地震"}
        b._category_sensors = {"振动": ["101"], "地震": ["201"]}
        b.metrics = {"vibration": {"feature": "DZJSD(xJsd)"},
                     "earthquake_load": {"feature": "DZJSD(xJsd)"}}
        b.resolve_metric_stat_detail = lambda m, stat, period: (None, {})
        b._is_excluded = lambda sid: False
        return b

    def test_extreme_category_isolation(self):
        b = self._bridge()
        period = self._period()
        pos_entries = {
            "振动位置": {"测点1": {"统计": {
                "最大值": 1.68, "最小值": -9.57, "差值": 11.25},
                "传感器编号": "101"}},
            "地震位置": {"测点1": {"统计": {
                "最大值": 5602250.0, "最小值": -0.1,
                "差值": 5602250.1}, "传感器编号": "201"}},
        }
        gs = {"最大值": 5602250.0, "最小值": -0.1,
              "差值": 5602250.1}
        v, loc = b._metric_extreme("vibration", "max", "最大值",
                                   "最大值位置", period, gs, pos_entries)
        self.assertAlmostEqual(v, 1.68, places=6)
        self.assertEqual(loc, "振动位置")
        # 地震类别唯一的传感器极值异常且无同族有效值：
        # 严格类别模式下不得回退到共享特征码的 5602250 全桥统计
        v2, _ = b._metric_extreme("earthquake_load", "max", "最大值",
                                  "最大值位置", period, gs, pos_entries)
        self.assertIsNone(v2)

    def _period(self):
        return {"start": dt.date(2026, 7, 1),
                "end": dt.date(2026, 9, 30),
                "label": "2026.7~9"}


class ChartCleaningGranularityTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location(
            "bcl_clean",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))),
                "preprocess", "scripts", "build_chart_library.py"))
        cls.bcl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.bcl)

    def test_granularity_mapping(self):
        g = self.bcl.feature_granularity
        self.assertEqual(g("DZJSD(yJsd)"), "second")
        self.assertEqual(g("SZJSD(xJsd)"), "second")
        self.assertEqual(g("FSFX2(spfs)"), "10min")
        self.assertEqual(g("WD(temp)"), "hour")

    def test_budget_scales_with_granularity(self):
        spk, dist, rm = self.bcl.granularity_cleaning_budget(
            "DZJSD(yJsd)", 5, 5, 5)
        self.assertGreaterEqual(spk, 200)
        spk2, _, _ = self.bcl.granularity_cleaning_budget(
            "FSFX2(spfs)", 5, 5, 5)
        self.assertGreaterEqual(spk2, 20)

    def test_gross_spike_removed_even_when_ratio_low(self):
        # 90% 数据在量程外但未到 10 倍量程（如量程漂移），10% 是
        # 1e7 超量级毛刺：命中率门槛会失效，但毛刺必须无条件剔除
        values = [50000.0] * 90 + [1e7] * 10
        times = [dt.datetime(2026, 7, 1) + dt.timedelta(hours=i)
                 for i in range(len(values))]
        out, recs, ix, rx = self.bcl.clean_series_value(
            times, values, "t", spike_k=5.0, hour_level=True,
            vrange=(-10000.0, 10000.0), max_spikes=200,
            max_dist_outliers=100, max_total_removals=100)
        self.assertLess(max(out), 1e6)


class PerfFastPathEquivalenceTest(unittest.TestCase):
    """性能优化（pandas 解析、向量化掩码）必须与原来的逐点实现等价。"""

    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location(
            "bcl_perf",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))),
                "preprocess", "scripts", "build_chart_library.py"))
        cls.bcl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.bcl)

    _CSV = "\n".join([
        "bucket_start,count,mean,min,max,sum,std,median",
        "2026-07-01T00:00:00,3600,1.5,1.0,2.0,0,0,",
        "2026-07-01T01:00:00,3600,-0.5,-1.0,0.0,0,0,",
        "坏行",
        "2026-07-01T02:00:00,x,1.0,0.5,1.5,0,0,",
        "2026-07-01T03:00:00,3600,nan,0.5,1.5,0,0,",
        "2026-07-01T04:00:00,0,1.0,0.5,1.5,0,0,",
        "2026-07-01T05:00:00,3600,3.5,3.0,4.0,0,0,",
    ])

    def test_two_read_paths_agree(self):
        """pandas 快速解析与纯 python 逐行解析结果必须完全一致。"""
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "2026-07-01.csv")
            with open(p, "w", encoding="utf-8") as f:
                f.write(self._CSV)
            slow = self.bcl.read_daily_file(p)          # 小文件 → 逐行
            try:
                fast = self.bcl._read_daily_file_pandas(p)
            except Exception as exc:                    # noqa: BLE001
                self.skipTest(f"环境无 pandas: {exc}")
            self.assertEqual(slow[:5], fast[:5])
            self.assertEqual(slow[5], fast[5])
            self.assertEqual(len(slow[0]), 3)

    def test_detect_zero_runs_matches_naive(self):
        times = [dt.datetime(2026, 7, 1) + dt.timedelta(hours=i)
                 for i in range(80)]
        for vals in ([0.0] * 30 + [1.0] * 50,
                     [1.0] * 5 + [0.0] * 28 + [2.0] * 47,
                     [0.0] * 80):
            got = self.bcl.detect_zero_runs(times, vals, min_hours=24.0)
            # 逐点参考实现（原逻辑）
            want, start = [], None
            for i, v in enumerate(vals):
                is_zero = abs(float(v)) <= 1e-9
                if is_zero and start is None:
                    start = i
                elif not is_zero and start is not None:
                    dur = (times[i - 1] - times[start]).total_seconds() / 3600.0
                    if dur > 24.0:
                        want.append((times[start], times[i - 1], round(dur, 1)))
                    start = None
            if start is not None:
                dur = (times[-1] - times[start]).total_seconds() / 3600.0
                if dur > 24.0:
                    want.append((times[start], times[-1], round(dur, 1)))
            self.assertEqual(
                [(dt.datetime.strptime(r["起始时间"], "%Y-%m-%d %H:%M"),
                  dt.datetime.strptime(r["结束时间"], "%Y-%m-%d %H:%M"),
                  r["持续小时数"]) for r in got], want)

    def test_mask_zero_run_hours_matches_naive(self):
        times = [dt.datetime(2026, 7, 1) + dt.timedelta(hours=i)
                 for i in range(60)]
        means = [float(i) for i in range(60)]
        maxs = [v + 1 for v in means]
        mins = [v - 1 for v in means]
        runs = [{"起始时间": "2026-07-01 10:00",
                 "结束时间": "2026-07-01 20:00"},
                {"起始时间": "2026-07-02 05:00",
                 "结束时间": "2026-07-02 07:00"}]
        spans = [(dt.datetime.strptime(r["起始时间"], "%Y-%m-%d %H:%M"),
                  dt.datetime.strptime(r["结束时间"], "%Y-%m-%d %H:%M"))
                 for r in runs]
        want = [i for i, h in enumerate(times)
                if not any(t0 <= h <= t1 for t0, t1 in spans)]
        got = self.bcl._mask_zero_run_hours(times, means, maxs, mins, runs)[0]
        self.assertEqual(got, [times[i] for i in want])

    def test_monthly_floor_array_matches_scalar(self):
        times = [dt.datetime(2026, 3, 30) + dt.timedelta(days=i)
                 for i in range(10)]
        got = self.bcl._monthly_floor_array(times)
        want = [self.bcl.seasonal_min_for("WD(temp)", t) for t in times]
        self.assertEqual(got.tolist(), want)


class SummaryMissThresholdTest(unittest.TestCase):
    """方案第 5 节：结论段只说明缺失 >7 天（168h）的时段。"""

    def test_default_is_seven_days_and_cfg_overrides(self):
        b = BridgeData.__new__(BridgeData)
        b.cfg = {}
        self.assertEqual(b._summary_miss_threshold(), 168.0)
        b.cfg = {"summary_miss_hours": 72}
        self.assertEqual(b._summary_miss_threshold(), 72.0)
        b.cfg = {"summary_miss_hours": None}
        self.assertEqual(b._summary_miss_threshold(), 168.0)


class RobustStatFaultTest(unittest.TestCase):
    """极值故障不应把均值/中位数这类稳健统计量也打成“—”。"""

    def _stats(self, **kw):
        base = {"平均值": 12.47, "中位数": 12.4, "均方根值": 13.0,
                "最大值": 51.19, "最小值": -21.90, "差值": 73.08}
        base.update(kw)
        return base

    def test_extreme_fault_keeps_average(self):
        f = self._stats()
        # 最低温 -21.9℃、差值 73℃ 属故障 → 极值判失效
        self.assertTrue(BridgeData._gross_stat_fault(f, "WD(temp)", "max"))
        self.assertTrue(BridgeData._gross_stat_fault(f, "WD(temp)", "min"))
        # 但平均 12.47℃ 正常 → 不能填“—”
        for stat in ("avg", "平均温度", "median", "rms"):
            self.assertFalse(BridgeData._gross_stat_fault(f, "WD(temp)", stat),
                             stat)

    def test_impossible_average_still_faulty(self):
        self.assertTrue(BridgeData._gross_stat_fault(
            self._stats(平均值=500.0), "WD(temp)", "avg"))
        self.assertTrue(BridgeData._gross_stat_fault(
            self._stats(平均值=-40.0), "WD(temp)", "avg"))

    def test_contradictory_average_still_faulty(self):
        """平均值不在 [最小值, 最大值] 内 → 口径自相矛盾，仍判失效。"""
        self.assertTrue(BridgeData._gross_stat_fault(
            self._stats(平均值=99.0), "WD(temp)", "avg"))

    def test_humidity_same_rule(self):
        # 极值离谱但均值本身合理（79.7%）→ 极值判失效、均值照常给
        f = {"平均值": 79.7, "最大值": 525843000.0, "最小值": 18.06,
             "差值": 525842980.0}
        self.assertTrue(BridgeData._gross_stat_fault(f, "WSD(rh)", "max"))
        self.assertTrue(BridgeData._gross_stat_fault(f, "WSD(rh)", "range"))
        self.assertFalse(BridgeData._gross_stat_fault(f, "WSD(rh)", "avg"))
        # 均值自己就超量程（>100%）→ 该行整体失效
        f2 = {"平均值": 114.56, "最大值": 525843000.0, "最小值": 19.63,
              "差值": 525842980.0}
        self.assertTrue(BridgeData._gross_stat_fault(f2, "WSD(rh)", "avg"))

    def test_header_style_stat_names_resolve(self):
        """表头式列名（平均温度/最高温度/最低温度/最大温差）要能归一化。"""
        from report_agent.bridge_source import _canon_stat
        self.assertEqual(_canon_stat("平均温度"), "avg")
        self.assertEqual(_canon_stat("平均湿度"), "avg")
        self.assertEqual(_canon_stat("最高温度"), "max")
        self.assertEqual(_canon_stat("最低温度"), "min")
        self.assertEqual(_canon_stat("最大温差"), "range")
        self.assertEqual(_canon_stat("均值"), "avg")
        self.assertEqual(_canon_stat("绝对值最大"), "abs_max")


class MissSevereThresholdTest(unittest.TestCase):
    """缺失严重口径：季度/月度按“缺失合计≥7天(168h)”，年度按>30天。"""

    @classmethod
    def setUpClass(cls):
        scripts = os.path.join(os.path.dirname(os.path.dirname(
            os.path.abspath(__file__))), "preprocess", "scripts")
        spec = importlib.util.spec_from_file_location(
            "bq_test", os.path.join(scripts, "build_quarterly_stats.py"))
        cls.bq = importlib.util.module_from_spec(spec)
        import sys as _sys
        if scripts not in _sys.path:
            _sys.path.insert(0, scripts)
        spec.loader.exec_module(cls.bq)

    def test_quarterly_uses_seven_days(self):
        f = self.bq.is_missing_severe
        self.assertFalse(f({"缺失小时数": 24, "缺失天数": 1}, "quarterly"))
        self.assertFalse(f({"缺失小时数": 72, "缺失天数": 3}, "quarterly"))
        self.assertFalse(f({"缺失小时数": 167, "缺失天数": 6}, "quarterly"))
        self.assertTrue(f({"缺失小时数": 168, "缺失天数": 7}, "quarterly"))
        self.assertTrue(f({"缺失小时数": 0, "缺失天数": 8}, "quarterly"))

    def test_yearly_uses_one_month(self):
        f = self.bq.is_missing_severe
        self.assertFalse(f({"缺失小时数": 168, "缺失天数": 7}, "yearly"))
        self.assertTrue(f({"缺失小时数": 900, "缺失天数": 31}, "yearly"))

    def test_report_missing_list_copies_quarterly_summary(self):
        """缺失清单以季度总结为准，并按“≥7天/168h”复核（旧总结是 72h 口径）。"""
        b = BridgeData({})
        # 绕过“指标类别隔离”（单元测试里没有传感器对照表）
        b._filter_pos_entries = lambda metric, pe: (pe, True)
        pe = {"甲位置": {"测点1": {"统计": {"缺失小时数": 200,
                                        "缺失天数": 8}}},
              "乙位置": {"测点1": {"统计": {"缺失小时数": 72,
                                        "缺失天数": 3}}}}
        gs = {"数据缺失严重的传感器位置": ["甲位置", "乙位置",
                                        "丙位置（无统计）"]}
        _zero, _seg, miss = b._fault_positions(gs, pe, "WD(temp)",
                                              "structure_temperature",
                                              {"label": "2026.7~9"})
        # 甲(≥168h) 保留；乙(72h，旧总结口径) 被复核掉；查不到的按总结保留
        self.assertIn("甲位置", miss)
        self.assertNotIn("乙位置", miss)
        self.assertIn("丙位置（无统计）", miss)

    def test_missing_label_says_over_seven_days(self):
        b = BridgeData({})
        self.assertIn("超过7天", b._miss_label(False))
        self.assertIn("168h", b._miss_label(False))
        self.assertIn("一个月", b._miss_label(True))


class SensorMapParseTest(unittest.TestCase):
    """《五座桥测点编号表格.docx》不会变 → 解析结果必须稳定、可复现。

    这里把“解析基线”钉住：传感器总数、各桥数量、测点/表格映射的结构与
    内容摘要。解析器一旦漏表（历史上就漏过 结构温度监测表，而报告填表
    依赖它）或文档被改动，测试会立刻失败。
    """

    # 解析基线摘要（文档/解析器没变时应保持不变；故意调整需同步更新）
    BASELINE_DIGEST = ("d4108adf5ec7cb25b251c0066198993"
                       "5e545c5112e52215050038737c089f04a")
    BRIDGE_COUNTS = {"湘江特大桥": 185, "洣水河特大桥": 213, "矮寨大桥": 286,
                     "赤石大桥": 223, "洞庭湖大桥": 250}

    @classmethod
    def setUpClass(cls):
        import hashlib
        import json as _json
        from collections import Counter
        cls.hashlib, cls.json, cls.Counter = hashlib, _json, Counter
        docx = os.path.join(os.path.dirname(os.path.dirname(
            os.path.abspath(__file__))), "inputs",
            "五座桥测点编号表格.docx")
        if not os.path.isfile(docx):
            raise unittest.SkipTest("缺少测点编号表 docx")
        spec = importlib.util.spec_from_file_location(
            "psm_test",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))),
                "preprocess", "scripts", "parse_sensor_map.py"))
        cls.m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.m)
        cls.sensors = cls.m.parse_docx(docx)

    def test_sensor_counts_pinned(self):
        self.assertEqual(len(self.sensors), sum(self.BRIDGE_COUNTS.values()))
        got = self.Counter(self.m.bridge_of(i) for i in self.sensors.values())
        for bridge, n in self.BRIDGE_COUNTS.items():
            self.assertEqual(got.get(bridge), n, bridge)

    def test_position_and_table_map_digest_stable(self):
        blob = {b: {"测点映射": self.m.build_position_map(self.sensors, b),
                    "表格映射": self.m.build_table_map(self.sensors, b, {})}
                for b in self.m.BRIDGE_ORDER}
        raw = self.json.dumps(blob, ensure_ascii=False, sort_keys=True)
        self.assertEqual(
            self.hashlib.sha256(raw.encode("utf-8")).hexdigest(),
            self.BASELINE_DIGEST,
            "解析结果与基线不一致：解析器改动或 docx 被替换，"
            "请核对差异（确认无误再更新 BASELINE_DIGEST）")

    def test_struct_temp_position_map_present(self):
        """结构温度监测表必须在测点映射里（报告按它定位行↔传感器）。"""
        for bridge in self.m.BRIDGE_ORDER:
            pm = self.m.build_position_map(self.sensors, bridge)
            self.assertIn("结构温度监测表", pm, bridge)
        plans = self.m.build_position_map(
            self.sensors, "洞庭湖大桥")["结构温度监测表"]
        by_pos = {p["断面位置"]: p["测点"] for p in plans}
        pts = by_pos.get("君山侧塔梁交接处钢桁梁") or {}
        self.assertEqual(len(pts), 20)
        self.assertEqual(pts.get("测点1"), "3268")

    def test_loss_of_detects_dropped_content(self):
        """写入前的护栏：新结果少表/少断面位置时必须报出来。"""
        old = {"测点映射": {"结构温度监测表": [
                  {"断面位置": "君山侧塔梁交接处钢桁梁", "测点": {}}]},
               "传感器": {"3268": {}}}
        new = {"测点映射": {"结构应变监测表": []}}
        losses = self.m.loss_of(old, new)
        self.assertTrue(any("结构温度监测表" in x for x in losses), losses)
        self.assertTrue(any("传感器编号少" in x for x in losses), losses)
        # 只新增（不丢）不算退步；少了表要报出来
        base = {"测点映射": {"结构应变监测表": []}}
        more = {"测点映射": {"结构应变监测表": [],
                             "结构温度监测表": []}}
        self.assertEqual(self.m.loss_of(base, more), [])
        self.assertTrue(self.m.loss_of(more, base))


class PreprocessPlanTest(unittest.TestCase):
    """前端“缺什么补什么”：统计值在、图库缺 → 只补图库（+季度/年度总结）。"""

    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location(
            "web_app",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))), "web", "app.py"))
        cls.app = importlib.util.module_from_spec(spec)
        try:
            spec.loader.exec_module(cls.app)
        except Exception as exc:  # noqa: BLE001
            raise unittest.SkipTest(f"web.app 不可导入: {exc}")

    @staticmethod
    def _make(tmp, charts: bool, stats: bool):
        charts_dir = os.path.join(tmp, "图库_2026.7~9", "测试桥")
        stats_dir = os.path.join(tmp, "统计值_2026.7~9", "测试桥")
        if charts:
            d = os.path.join(charts_dir, "某位置", "WD(temp)")
            os.makedirs(d, exist_ok=True)
            with open(os.path.join(d, "时间序列图.png"), "wb") as f:
                f.write(b"x")
        if stats:
            d = os.path.join(stats_dir, "位置统计", "某位置")
            os.makedirs(d, exist_ok=True)
            with open(os.path.join(d, "WD(temp).json"), "w",
                      encoding="utf-8") as f:
                f.write("{}")
        return charts_dir, stats_dir

    def test_plan_by_missing_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            for charts, stats, daily, want in (
                    (True, True, True, "skipped"),
                    (False, True, True, "charts_only"),
                    (True, False, True, "stats_only"),
                    (False, False, True, "full"),
                    (False, True, False, "full"),   # 日级也要补 → 走完整流程
            ):
                charts_dir, stats_dir = self._make(
                    os.path.join(tmp, f"{charts}{stats}{daily}"), charts, stats)
                self.assertEqual(
                    self.app._preprocess_plan(charts_dir, stats_dir, daily),
                    want, (charts, stats, daily))
                self.assertIn(want, self.app.PREPROCESS_PLAN_TEXT)


class ChartBeautifyTest(unittest.TestCase):
    """方案第 2/3/4/6 节：图上不画文字、子图等宽等高、图例精简、Z-score。"""

    @classmethod
    def setUpClass(cls):
        import matplotlib
        matplotlib.use("Agg")
        spec = importlib.util.spec_from_file_location(
            "bcl_beautify",
            os.path.join(os.path.dirname(os.path.dirname(
                os.path.abspath(__file__))),
                "preprocess", "scripts", "build_chart_library.py"))
        cls.bcl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.bcl)

    @staticmethod
    def _temp_series():
        base = dt.datetime(2026, 7, 1)
        hours = [base + dt.timedelta(hours=i) for i in range(240)]
        series = []
        for sid, off in (("434", 0.0), ("435", 3.0), ("436", -2.0)):
            means = [20.0 + off + (i % 24) / 24.0 for i in range(240)]
            series.append({
                "label": sid, "feature": "WD(temp)", "sensor": sid,
                "hours": hours, "means": means,
                "spike_pts": [(hours[10], means[10])],
                "range_pts": [(hours[11], means[11])],
                "gaps": [{"起始时间": "2026-07-03 00:00",
                          "结束时间": "2026-07-05 00:00",
                          "缺失小时数": 48}],
                "records": [],
                "shifts": [{"起始时间": "2026-07-06 00:00",
                            "结束时间": "2026-07-08 00:00",
                            "方向": "偏高"}],
            })
        return series

    def test_group_chart_has_no_text_and_equal_panels(self):
        """时间序列合并图：图上无文字标注，子图等宽等高，画布不畸变。"""
        plt = self.bcl.plt
        orig_close = plt.close
        plt.close = lambda *a, **k: None
        try:
            with tempfile.TemporaryDirectory() as tmp:
                out = os.path.join(tmp, "时间序列图.png")
                self.bcl.plot_group_time_series(
                    "测试位置", "WD(temp)", self._temp_series(), out)
                self.assertTrue(os.path.isfile(out))
                fig = plt.figure(plt.get_fignums()[-1])
                axes = [a for a in fig.axes if a.get_visible()]
                self.assertGreaterEqual(len(axes), 3)
                # 图上不得残留“时间段/偏高偏低/可能故障”文字标注
                for ax in axes:
                    self.assertEqual(len(ax.texts), 0)
                # 同一张图内所有子图等宽等高
                boxes = [a.get_position() for a in axes]
                for b in boxes[1:]:
                    self.assertAlmostEqual(b.width, boxes[0].width, places=4)
                    self.assertAlmostEqual(b.height, boxes[0].height,
                                           places=4)
                # 无右侧留白：子图占满画布宽度（旧实现收缩到 ~0.62）
                self.assertGreater(max(b.x1 for b in boxes), 0.9)
                # 单曲线图例只留状态项，最多 4 条，一行放得下
                self.assertTrue(fig.legends)
                legend_labels = [t.get_text()
                                 for t in fig.legends[0].get_texts()]
                self.assertLessEqual(len(legend_labels), 4)
                self.assertNotIn("434", legend_labels)
                from PIL import Image
                with Image.open(out) as im:
                    w, h = im.size
                self.assertLess(w / float(h), 8.0)
        finally:
            plt.close = orig_close
            plt.close("all")

    def test_legend_matches_drawn_bands(self):
        """图上画了色带，图例就必须有对应条目（不能“有彩带没图例”）。"""
        import datetime as _dt
        plt = self.bcl.plt
        base = _dt.datetime(2026, 7, 1)
        hours = [base + _dt.timedelta(hours=i) for i in range(24 * 20)]
        means = []
        for i, h in enumerate(hours):
            means.append(0.0 if 100 <= i < 160
                         else 20.0 + (h.hour % 24) * 0.2)
        series = [{
            "label": "434", "feature": "WD(temp)", "sensor": "434",
            "hours": hours, "means": means,
            "spike_pts": [(hours[20], 20.0)],
            "range_pts": [(hours[30], 20.0)],
            "gaps": [{"起始时间": "2026-07-05 00:00",
                      "结束时间": "2026-07-06 12:00", "缺失小时数": 36}],
            "records": [],
            "shifts": [{"起始时间": "2026-07-08 00:00",
                        "结束时间": "2026-07-10 00:00", "方向": "偏高"},
                       {"起始时间": "2026-07-12 00:00",
                        "结束时间": "2026-07-14 00:00", "方向": "偏低"}],
        }]
        orig_close = plt.close
        plt.close = lambda *a, **k: None
        try:
            with tempfile.TemporaryDirectory() as tmp:
                out = os.path.join(tmp, "时间序列图.png")
                self.bcl.plot_group_time_series("测试位置", "WD(temp)",
                                                series, out)
                fig = plt.figure(plt.get_fignums()[-1])
                ax = [a for a in fig.axes if a.get_visible()][0]
                bands = [p for p in ax.patches]          # axvspan 色带
                labels = ([t.get_text() for t in fig.legends[0].get_texts()]
                          if fig.legends else [])
                self.assertEqual(len(bands), 4)          # 橙/红/绿/紫
                for lb in ("数据缺失(已插值填充)", "长时间偏高",
                           "长时间偏低", "可能故障(恒0超过24h)",
                           "已替换尖峰(统计)", "已剔除异常值(范围外)"):
                    self.assertIn(lb, labels)
                # 图例必须完整落在画布内（不能被裁掉）
                fig.canvas.draw()
                ext = fig.legends[0].get_window_extent(
                    fig.canvas.get_renderer())
                h_px = fig.get_size_inches()[1] * fig.dpi
                self.assertGreaterEqual(ext.y0, -1.0)
                self.assertLessEqual(ext.y1, h_px + 1.0)
        finally:
            plt.close = orig_close
            plt.close("all")

    def test_trim_legend_drops_ids_for_single_curve(self):
        labels = ["434", "长时间偏高", "可能故障(恒0超过24h)"]
        handles = [1, 2, 3]
        kept_l, _ = self.bcl._trim_legend_items(labels, handles, 1)
        self.assertNotIn("434", kept_l)
        self.assertIn("长时间偏高", kept_l)
        multi_l, _ = self.bcl._trim_legend_items(labels, handles, 3, True)
        self.assertIn("434", multi_l)

    def test_stat_window_by_granularity(self):
        """统计值清洗窗口按特征粒度取：风速 24h、加速度 1h(=原逻辑)、
        温度等保持原逻辑(0=全局)。"""
        w = self.bcl._stat_window_for
        self.assertEqual(w("FSFX2(spfs)"), 24)
        self.assertEqual(w("DZJSD(yJsd)"), 0)     # 统计序列本身已是小时级
        self.assertEqual(w("SZJSD(xJsd)"), 0)
        self.assertEqual(w("WD(temp)"), 0)
        self.assertEqual(w("FSFX2(spfs)", 6), 6)  # 显式窗口优先
        self.assertEqual(w("FSFX2(spfs)", -1), 0)  # 负数=强制全局

    def test_no_zscore_logic_left(self):
        """Z-score 整套逻辑已删除（含绘图与命令行），不得再被引入。"""
        import inspect
        src = inspect.getsource(self.bcl).lower()
        self.assertNotIn("zscore", src)
        for fn in (self.bcl.read_clean_hourly_means,
                   self.bcl._build_merged_series):
            params = inspect.signature(fn).parameters
            self.assertNotIn("zscore_k", params)
            self.assertNotIn("zscore_window", params)
            self.assertNotIn("stat_window", params)
        # 统计值清洗入口只保留 feature/stat_window 两个统计侧参数
        params = inspect.signature(self.bcl.clean_series_value).parameters
        self.assertIn("feature", params)
        self.assertIn("stat_window", params)

    def test_seasonal_temperature_floor_removes_impossible_cold(self):
        """7 月不可能低于 0℃：少数不合理的负值按异常剔除。"""
        times = [dt.datetime(2026, 7, 1) + dt.timedelta(hours=i)
                 for i in range(240)]
        vals = [30.0 + (i % 24) * 0.1 for i in range(240)]
        vals[100] = -20.0
        vals[150] = -5.0
        out, recs, _spike, rng = self.bcl.clean_series_value(
            times, vals, "t", spike_k=0.0, hour_level=False,
            vrange=(-30.0, 70.0), feature="WD(temp)", dist_k=0.0)
        self.assertIn(100, rng)
        self.assertIn(150, rng)
        self.assertGreater(out[100], 0.0)
        self.assertTrue(any("季节" in str(r.get("说明", ""))
                            for r in recs))

    def test_seasonal_floor_not_applied_when_whole_series_below(self):
        """整段低于当月下限（量程不同/单位不同）时不硬过滤，避免清空序列。"""
        times = [dt.datetime(2026, 7, 1) + dt.timedelta(hours=i)
                 for i in range(240)]
        vals = [-40.0 - (i % 24) * 0.1 for i in range(240)]
        out, recs, _spike, rng = self.bcl.clean_series_value(
            times, vals, "t", spike_k=0.0, hour_level=False,
            feature="WD(temp)", dist_k=0.0, max_spikes=0,
            max_total_removals=0)
        self.assertEqual(rng, [])
        self.assertTrue(any("未硬过滤" in str(r.get("说明", ""))
                            for r in recs))

    def test_summer_zero_reading_removed(self):
        """夏季(7~9月)恰好 0.0℃ 的掉零读数也要按季节异常剔除——
        否则最小值会被打成 0、差值等于最大值（洣水河 636 那种情况）。"""
        times = [dt.datetime(2026, 8, 1) + dt.timedelta(hours=i)
                 for i in range(240)]
        vals = [25.0 + (i % 24) * 0.05 for i in range(240)]
        vals[100] = 0.0
        vals[101] = 0.0
        out, recs, _spike, rng = self.bcl.clean_series_value(
            times, vals, "t", spike_k=0.0, hour_level=False,
            vrange=(-30.0, 70.0), feature="WD(temp)", dist_k=0.0,
            max_spikes=0, max_total_removals=0)
        self.assertIn(100, rng)
        self.assertIn(101, rng)
        self.assertGreater(min(out), 1.0)
        self.assertTrue(any("季节" in str(r.get("说明", ""))
                            for r in recs))

    def test_stat_window_catches_local_spike_global_misses(self):
        """风速统计按 24h 局部窗口判定：日际变化大时全局带太宽，
        会漏掉局部阵风；窗口模式能发现并剔除。"""
        import math
        n = 24 * 90
        times = [dt.datetime(2026, 1, 1) + dt.timedelta(hours=i)
                 for i in range(n)]
        vals = [5.0 * math.sin(i / (24 * 30.0) * 2 * math.pi)
                for i in range(n)]
        vals[1000] += 30.0                      # 相对当日上下文异常的阵风
        kw = dict(spike_k=0.0, hour_level=False, max_spikes=1,
                  max_total_removals=1, dist_k=0.0)
        _o1, _r1, s1, x1 = self.bcl.clean_series_value(times, vals, "t", **kw)
        out2, recs2, _s2, x2 = self.bcl.clean_series_value(
            times, vals, "t", stat_window=24, feature="FSFX2(spfs)", **kw)
        self.assertEqual((s1, x1), ([], []))    # 全局逻辑漏检
        self.assertIn(1000, x2)                 # 24h 窗口命中
        self.assertLess(abs(out2[1000]), 10.0)  # 用局部基线替代
        self.assertTrue(any("24h 局部窗口" in str(r.get("说明", ""))
                            for r in recs2))


class ScatterFallbackTest(unittest.TestCase):
    def test_folder_name_order_tolerance(self):
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        name_dict = os.path.join(root, "preprocess", "传感器对照",
                                 "传感器名称对照", "洣水河特大桥.json")
        if not os.path.isfile(name_dict):
            self.skipTest("缺少洣水河名称对照表")
        with tempfile.TemporaryDirectory() as tmp:
            d = os.path.join(tmp, "左幅炎陵侧边跨跨中顶板")
            os.makedirs(d)
            png = os.path.join(d, "相关性_WD(temp)-YB(rsg).png")
            with open(png, "wb") as f:
                f.write(b"x")
            import json as _json
            cfgp = os.path.join(root, "config", "config_mishuihe.json")
            if not os.path.isfile(cfgp):
                self.skipTest("缺少洣水河配置")
            cfg = _json.load(open(cfgp, encoding="utf-8"))["bridge_data"]
            cfg = dict(cfg)
            cfg["sensor_map"] = os.path.join(
                root, "preprocess", "传感器对照",
                "传感器编号名称.json")
            cfg["name_dict"] = name_dict
            cfg["stats_dir"] = ""
            cfg["charts_dir"] = tmp
            b = BridgeData(cfg)
            b.load()
            info = b.resolve_chart_info(
                "strain_左幅炎陵侧边跨跨中顶板_scatter_57",
                "strain_左幅炎陵侧边跨跨中顶板_scatter_57")
            self.assertIsNotNone(info)
            self.assertTrue(info["path"].endswith(".png"))

    def test_static_keep_works_without_llm(self):
        analysis = {
            "texts": ["A(2,2)", "最大10min平均风速为11.93 m/s。"],
            "numbers": [
                {"value": "2", "paragraph": 0, "verdict": "replace",
                 "placeholder": None, "reasons": []},
                {"value": "11.93", "paragraph": 1, "verdict": "replace",
                 "placeholder": None, "reasons": []},
            ],
        }
        stats = arbitrate_numbers(analysis, {"enabled": False})
        self.assertFalse(stats["enabled"])
        self.assertGreaterEqual(stats["static_keep"], 1)
        self.assertEqual(analysis["numbers"][0]["verdict"], "keep")
        self.assertEqual(analysis["numbers"][1]["verdict"], "replace")


if __name__ == "__main__":
    unittest.main()
