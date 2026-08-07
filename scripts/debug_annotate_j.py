#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Debug: run annotate_docx with recognize (no LLM) and check J-table."""

import os
import sys
import re

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from report_agent.recognizer import recognize, annotate_docx, CELL_REF_RE

SRC = r"D:\五座桥数据分析文件\洞庭湖大桥.docx"
DST = r"D:\Code\data_analysis\outputs\debug_j_test.docx"

print("Running recognize (no LLM)...")
analysis = recognize(SRC, llm_cfg=None)

# Check J-table cell_refs in analysis
j_refs = [ct for ct in analysis.get("chart_texts", [])
          if ct.get("source") == "cell_ref" and ct.get("table_letter") == "J"]
j_replace = [c for c in j_refs if c.get("verdict") == "replace"]
print(f"J-table cell_refs: {len(j_refs)} total, {len(j_replace)} replace")

# Check num_targets overlap with J paragraphs
j_para_indices = set(ct["paragraph"] for ct in j_replace)
num_in_j = []
for n in analysis.get("numbers", []):
    if n.get("paragraph") in j_para_indices:
        num_in_j.append(n)
        print(f"  num_target: para={n['paragraph']}, pos={n['position']}, "
              f"value={n['value']}, verdict={n['verdict']}, placeholder={n.get('placeholder')}")

# Build num_targets as annotate_docx would
num_targets = {}
for n in analysis["numbers"]:
    if n["verdict"] != "replace":
        continue
    marker = n.get("placeholder")
    if marker is None:
        continue
    elif not marker.startswith("{{"):
        marker = f"{{{{{marker}}}}}"
    num_targets.setdefault(n["paragraph"], {})[n["position"]] = marker

j_num_replace = {pi: num_targets.get(pi, {}) for pi in j_para_indices}
print(f"\nnum_targets (verdict=replace) overlapping J paragraphs:")
for pi, targets in j_num_replace.items():
    if targets:
        print(f"  para {pi}: {targets}")

print("\nRunning annotate_docx...")
result = annotate_docx(SRC, DST, analysis=analysis)

print(f"\nannotate_docx result:")
for k, v in result.items():
    if k != "output":
        print(f"  {k}: {v}")

# Check output for remaining J( references
print("\n" + "=" * 60)
print("Checking output for unreplaced cell_refs...")
print("=" * 60)
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from report_agent.report_builder import iter_block_items

doc = Document(DST)
j_remaining = []
h_remaining = []
m_remaining = []
cell_wind = []
cell_cable = []
cell_rot = []

for item in iter_block_items(doc):
    if isinstance(item, (Paragraph,)):
        text = item.text
        if re.search(r"J\(\d+,\s*\d+\)", text): j_remaining.append(text[:60])
        if re.search(r"H\(\d+,\s*\d+\)", text): h_remaining.append(text[:60])
        if re.search(r"M\(\d+,\s*\d+\)", text): m_remaining.append(text[:60])
        if "{{cell.wind_speed" in text: cell_wind.append(text[:80])
        if "{{cell.cable_clamp" in text: cell_cable.append(text[:80])
        if "{{cell.rotation" in text: cell_rot.append(text[:80])
    elif isinstance(item, Table):
        for row in item.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text
                    if re.search(r"J\(\d+,\s*\d+\)", text): j_remaining.append(text[:60])
                    if re.search(r"H\(\d+,\s*\d+\)", text): h_remaining.append(text[:60])
                    if re.search(r"M\(\d+,\s*\d+\)", text): m_remaining.append(text[:60])
                    if "{{cell.wind_speed" in text: cell_wind.append(text[:80])
                    if "{{cell.cable_clamp" in text: cell_cable.append(text[:80])
                    if "{{cell.rotation" in text: cell_rot.append(text[:80])

print(f"\nRemaining cell_refs in output:")
print(f"  H: {len(h_remaining)} remaining, {len(cell_cable)} placeholders")
print(f"  M: {len(m_remaining)} remaining, {len(cell_rot)} placeholders")
print(f"  J: {len(j_remaining)} remaining, {len(cell_wind)} placeholders")

if j_remaining:
    print(f"\nJ remaining samples:")
    for t in j_remaining[:5]:
        print(f"  '{t}'")
if cell_wind:
    print(f"\nJ (wind_speed) cell placeholders:")
    for t in cell_wind[:5]:
        print(f"  '{t}'")

print("\nDone.")
